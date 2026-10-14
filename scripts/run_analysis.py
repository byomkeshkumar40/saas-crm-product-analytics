#!/usr/bin/env python3
"""Build the SQLite database from data/raw/*.csv, run sql/product_analysis.sql, export results.

Outputs
  data/processed/crm_analytics.db        SQLite database
  data/processed/kpi_tables/<query>.csv   one CSV per analysis query
  data/processed/powerbi/*.csv            Power BI model tables (dim_*, fact_events)
  data/processed/kpi_tables.xlsx          Excel workbook of the KPI tables (needs openpyxl)
"""
import csv
import re
import sqlite3
from datetime import date, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw"
PROC = ROOT / "data" / "processed"
DB_PATH = PROC / "crm_analytics.db"
KPI_DIR = PROC / "kpi_tables"
PBI_DIR = PROC / "powerbi"

TABLES = {  # load order respects foreign keys
    "accounts": ["account_id", "signup_date", "industry", "company_size", "subscription_plan", "region"],
    "users": ["user_id", "account_id", "created_at", "user_role"],
    "subscriptions": ["subscription_id", "account_id", "plan", "subscription_start", "subscription_status"],
    "events": ["event_id", "user_id", "account_id", "event_timestamp", "event_name", "feature_name",
               "session_id"],
}


def build_db():
    PROC.mkdir(parents=True, exist_ok=True)
    DB_PATH.unlink(missing_ok=True)
    con = sqlite3.connect(DB_PATH)
    con.executescript((ROOT / "sql" / "schema.sql").read_text())
    con.execute("PRAGMA foreign_keys = ON")
    for table, cols in TABLES.items():
        with open(RAW / f"{table}.csv", newline="", encoding="utf-8") as fh:
            rd = csv.DictReader(fh)
            rows = [tuple(r[c] for c in cols) for r in rd]
        con.executemany(f"INSERT INTO {table} ({','.join(cols)}) VALUES ({','.join('?' * len(cols))})", rows)
        print(f"loaded {table:<14} {len(rows):>7} rows")
    con.commit()
    fk = con.execute("PRAGMA foreign_key_check").fetchall()
    assert not fk, f"foreign key violations: {fk[:5]}"
    return con


def split_queries(sql_text):
    parts = re.split(r"^-- name: (\w+)\s*$", sql_text, flags=re.M)
    return list(zip(parts[1::2], parts[2::2]))


def write_csv(path, header, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as fh:
        wr = csv.writer(fh)
        wr.writerow(header)
        wr.writerows(rows)


def main():
    con = build_db()
    for d in (KPI_DIR, PBI_DIR):
        d.mkdir(parents=True, exist_ok=True)
        for f in d.glob("*.csv"):
            f.unlink()
    results = {}
    for name, body in split_queries((ROOT / "sql" / "product_analysis.sql").read_text()):
        cur = con.execute(body.strip())
        header = [c[0] for c in cur.description]
        rows = cur.fetchall()
        if name.startswith("model_"):
            out = PBI_DIR / ("fact_events.csv" if name == "model_fact_events"
                             else name.replace("model_", "", 1) + ".csv")
        else:
            out = KPI_DIR / f"{name}.csv"
            results[name] = (header, rows)
        write_csv(out, header, rows)
        print(f"{name:<34} {len(rows):>7} rows -> {out.relative_to(ROOT)}")

    # Date dimension covering the full observation window
    start, end = con.execute("SELECT obs_start, obs_end FROM analysis_params").fetchone()
    d0, d1 = date.fromisoformat(start), date.fromisoformat(end)
    dates = []
    d = d0
    while d <= d1:
        dates.append((d.isoformat(), d.year, d.month, d.strftime("%Y-%m"), d.strftime("%B"),
                      d.isocalendar()[1], d.isoweekday(), d.strftime("%a"),
                      "Weekend" if d.weekday() >= 5 else "Weekday"))
        d += timedelta(days=1)
    write_csv(PBI_DIR / "dim_date.csv",
              ["date", "year", "month_number", "year_month", "month_name", "iso_week",
               "iso_weekday", "weekday_name", "day_type"], dates)
    features = [("Lead Management", "Core", 1), ("Contact Management", "Core", 2),
                ("Deal Pipeline", "Core", 3), ("Email Integration", "Core", 4),
                ("Reporting", "Advanced", 5), ("Workflow Automation", "Advanced", 6)]
    write_csv(PBI_DIR / "dim_feature.csv", ["feature_name", "feature_tier", "sort_order"], features)
    print(f"dim_date / dim_feature written to {PBI_DIR.relative_to(ROOT)}")

    # Excel workbook of KPI tables
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        wb = Workbook()
        wb.remove(wb.active)
        for name, (header, rows) in results.items():
            if name == "q10_inactive_accounts":
                continue  # long account list stays CSV-only
            ws = wb.create_sheet(name[:31])
            ws.append(header)
            for c in ws[1]:
                c.font = Font(bold=True)
            for r in rows:
                ws.append(list(r))
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = min(
                    max(len(str(c.value)) if c.value is not None else 0 for c in col) + 2, 48)
            ws.freeze_panes = "A2"
        wb.save(PROC / "kpi_tables.xlsx")
        print("wrote data/processed/kpi_tables.xlsx")
    except ImportError:
        print("openpyxl not installed - skipping kpi_tables.xlsx (pip install -r requirements.txt)")
    con.close()


if __name__ == "__main__":
    main()
